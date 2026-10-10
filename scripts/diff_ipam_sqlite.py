#!/usr/bin/env python3
"""
Offline IPAM Diff Engine — NetBox Hub v4.3

Compares a multi-sheet Excel IP/VLAN inventory against a local SQLite vault
and emits three deterministic CSV reports:

  * ipam_matched.csv     — records present in both sources with matching site/VLAN/subnet
  * ipam_pending_add.csv — records found only in Excel (NetBox Universal Uploader format)
  * ipam_conflicts.csv   — records whose key exists in DB but site/VLAN/subnet differs

Zero cloud / LLM leakage. Pure local processing with openpyxl + sqlite3 + ipaddress.
"""

import argparse
import csv
import os
import sys
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
import ipaddress

# ── Path setup ──────────────────────────────────────────────────────────────

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
sys.path.insert(0, PROJECT_ROOT)


# ── ExcelIPAMParser ─────────────────────────────────────────────────────────

_HEADER_CANDIDATES: Dict[str, List[str]] = {
    "ip": ["ip", "ip address", "address", "ipv4", "host"],
    "subnet": ["subnet", "network", "cidr", "prefix"],
    "vlan_id": ["vlan id", "vid", "vlan", "vlan id#", "vlanid"],
    "vlan_name": ["vlan name", "name", "vlan description"],
    "description": ["description", "purpose", "desc", "notes", "comment"],
    "site": ["site", "location", "branch", "site name", "facility"],
}


def _normalise_header(h: str) -> str:
    return h.strip().lower().replace(" ", "_").replace("-", "_")


class ExcelIPAMParser:
    """Parse a multi-sheet Excel workbook into normalised IPAM records."""

    def __init__(self, excel_path: str):
        self.excel_path = os.path.abspath(excel_path)
        self.wb: Optional[openpyxl.Workbook] = None

    def _find_header_map(self, headers: List[str]) -> Dict[str, int]:
        """Map canonical column names to sheet column indices by inspecting the
        first row of headers against known candidate keywords.

        Both the sheet header tokens and the candidate keywords are normalised
        (lowercased, whitespace/hyphens collapsed to underscores) so that
        variants like ``VLAN ID``, ``Vid``, and ``vlan id`` all match ``vlan_id``.
        """
        mapping: Dict[str, int] = {}
        # Pre-normalise all candidate keywords for O(1) lookup per header token.
        _norm_candidates: Dict[str, str] = {}
        for canonical, candidates in _HEADER_CANDIDATES.items():
            for cand in candidates:
                _norm_candidates[_normalise_header(cand)] = canonical

        lower_headers = [_normalise_header(h) for h in headers]
        for idx, lh in enumerate(lower_headers):
            if lh in _norm_candidates and _norm_candidates[lh] not in mapping:
                mapping[_norm_candidates[lh]] = idx
        return mapping

    def _parse_row(
        self,
        row_vals: List[Any],
        header_map: Dict[str, int],
        default_site: str,
    ) -> Optional[Dict[str, str]]:
        """Extract and validate a single data row into a normalised record."""
        ip_raw = str(row_vals[header_map["ip"]]).strip() if "ip" in header_map else ""
        subnet_raw = str(row_vals[header_map["subnet"]]).strip() if "subnet" in header_map else ""
        vlan_id_raw = str(row_vals[header_map["vlan_id"]]).strip() if "vlan_id" in header_map else ""
        vlan_name_raw = str(row_vals[header_map["vlan_name"]]).strip() if "vlan_name" in header_map else ""
        desc_raw = str(row_vals[header_map["description"]]).strip() if "description" in header_map else ""
        site_raw = str(row_vals[header_map["site"]]).strip() if "site" in header_map else default_site

        # Normalise subnet — handle dot-notation like "10.113.64." with trailing mask
        subnet_clean = self._normalise_subnet(subnet_raw)
        if not subnet_clean:
            return None

        # Validate IP
        ip_clean = self._normalise_ip(ip_raw)
        if not ip_clean:
            return None

        # Skip noise / empty rows
        if all(v in ("", "nan", "NaN", "None", "null", "NULL") for v in [
            ip_raw, subnet_raw, vlan_id_raw, vlan_name_raw, desc_raw, site_raw
        ]):
            return None

        return {
            "ip_address": ip_clean,
            "subnet": subnet_clean,
            "vlan_id": vlan_id_raw if vlan_id_raw not in ("", "nan", "NaN") else "",
            "vlan_name": vlan_name_raw,
            "site": site_raw if site_raw not in ("", "nan", "NaN") else default_site,
            "description": desc_raw,
            "status": "Active",
        }

    @staticmethod
    def _normalise_ip(raw: str) -> str:
        """Return a valid IPv4 address string or empty string if invalid."""
        if not raw or raw.lower() in ("nan", "none", "null", ""):
            return ""
        try:
            addr = ipaddress.IPv4Address(raw.strip())
            return str(addr)
        except (ipaddress.AddressValueError, ValueError):
            return ""

    @staticmethod
    def _normalise_subnet(raw: str) -> str:
        """Return a valid CIDR string or empty string if invalid.

        Accepts full CIDR (``10.0.0.0/24``), bare IP addresses (treated as
        ``/32``), and dot-notation shorthand where a fifth octet represents
        the prefix length (``10.113.64.0.24`` → ``10.113.64.0/24``).
        """
        if not raw or raw.lower() in ("nan", "none", "null", ""):
            return ""
        s = raw.strip()
        # Handle dot-notation shorthand: "10.113.64.0.24" → "10.113.64.0/24"
        import re
        dot_match = re.match(
            r"^((?:\d{1,3}\.){3}\d{1,3})\.(\d{1,2})$", s
        )
        if dot_match:
            s = f"{dot_match.group(1)}/{dot_match.group(2)}"
        try:
            net = ipaddress.IPv4Network(s, strict=False)
            return str(net)
        except (ipaddress.AddressValueError, ValueError):
            return ""

    def parse(self) -> List[Dict[str, str]]:
        """Open the workbook read-only and return all valid records across sheets."""
        self.wb = openpyxl.load_workbook(
            self.excel_path, read_only=True, data_only=True
        )
        all_records: List[Dict[str, str]] = []
        seen_keys: Set[Tuple[str, str]] = set()

        for sheet_name in self.wb.sheetnames:
            ws = self.wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)

            # Discover header row within first 5 rows
            header_row_idx: Optional[int] = None
            header_values: List[str] = []
            for i, row in enumerate(rows_iter):
                if i >= 5:
                    break
                if row and any(cell is not None for cell in row):
                    header_values = [str(c).strip() if c is not None else "" for c in row]
                    header_row_idx = i
                    break

            if header_row_idx is None:
                continue

            header_map = self._find_header_map(header_values)

            # If no IP/subnet column found, skip
            if "ip" not in header_map and "subnet" not in header_map:
                continue

            default_site = sheet_name

            # rows_iter is a generator that has already consumed up to and
            # including the header row; every remaining row is a data row.
            for row in rows_iter:
                if row is None:
                    continue
                vals = list(row)
                # Pad to header length if needed
                while len(vals) < len(header_values):
                    vals.append("")
                rec = self._parse_row(vals, header_map, default_site)
                if rec is None:
                    continue
                # Deduplicate by (ip_address, subnet)
                key = (rec["ip_address"], rec["subnet"])
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                all_records.append(rec)

        self.wb.close()
        return all_records


# ── LocalSQLiteManager ──────────────────────────────────────────────────────

class LocalSQLiteManager:
    """Connect to the local sanitizer vault and manage ipam_inventory records."""

    def __init__(self, db_path: str):
        self.db_path = os.path.abspath(db_path)
        self._conn = __import__("sqlite3").connect(self.db_path, check_same_thread=False)
        self._conn.row_factory = __import__("sqlite3").Row
        self._ensure_table()

    def _ensure_table(self) -> None:
        self._conn.execute("""
            CREATE TABLE IF NOT EXISTS ipam_inventory (
                id        INTEGER PRIMARY KEY AUTOINCREMENT,
                ip_address TEXT UNIQUE,
                subnet    TEXT,
                vlan_id   TEXT,
                vlan_name TEXT,
                site      TEXT,
                description TEXT,
                status    TEXT DEFAULT 'Active'
            )
        """)
        self._conn.commit()

    def get_all_records(self) -> List[Dict[str, Any]]:
        """Return every row in ipam_inventory."""
        rows = self._conn.execute("SELECT * FROM ipam_inventory").fetchall()
        return [dict(r) for r in rows]

    def get_by_ip(self, ip: str) -> Optional[Dict[str, Any]]:
        row = self._conn.execute(
            "SELECT * FROM ipam_inventory WHERE ip_address = ?", (ip,)
        ).fetchone()
        return dict(row) if row else None

    def get_by_subnet(self, subnet: str) -> Optional[Dict[str, Any]]:
        row = self._conn.execute(
            "SELECT * FROM ipam_inventory WHERE subnet = ?", (subnet,)
        ).fetchone()
        return dict(row) if row else None

    def close(self) -> None:
        self._conn.close()


# ── Deterministic Comparison Logic ─────────────────────────────────────────

def classify_records(
    excel_records: List[Dict[str, str]],
    db_manager: LocalSQLiteManager,
) -> Tuple[List[Dict[str, str]], List[Dict[str, str]], List[Dict[str, str]]]:
    """Classify Excel records into matched / pending-add / conflicts."""
    matched: List[Dict[str, str]] = []
    pending_add: List[Dict[str, str]] = []
    conflicts: List[Dict[str, str]] = []

    for rec in excel_records:
        ip = rec["ip_address"]
        subnet = rec["subnet"]
        db_rec = db_manager.get_by_ip(ip) or db_manager.get_by_subnet(subnet)

        if db_rec is None:
            # Not in DB at all → pending add
            pending_add.append({
                "ip_address": ip,
                "subnet": subnet,
                "vlan_id": rec["vlan_id"],
                "vlan_name": rec["vlan_name"],
                "site": rec["site"],
                "description": rec["description"],
                "status": "Active",
            })
        else:
            # Key exists in DB — check for mismatch
            db_site = str(db_rec.get("site") or "").strip()
            db_vlan_id = str(db_rec.get("vlan_id") or "").strip()
            db_subnet = str(db_rec.get("subnet") or "").strip()
            excel_site = rec["site"].strip()
            excel_vlan_id = rec["vlan_id"].strip()

            site_mismatch = db_site.lower() != excel_site.lower()
            vlan_mismatch = db_vlan_id != excel_vlan_id
            subnet_mismatch = db_subnet != subnet

            if site_mismatch or vlan_mismatch or subnet_mismatch:
                reasons = []
                if site_mismatch:
                    reasons.append(
                        f"site Excel({excel_site}) vs DB({db_site})"
                    )
                if vlan_mismatch:
                    reasons.append(
                        f"vlan_id Excel({excel_vlan_id}) vs DB({db_vlan_id})"
                    )
                if subnet_mismatch:
                    reasons.append(
                        f"subnet Excel({subnet}) vs DB({db_subnet})"
                    )
                conflict_rec = {
                    **rec,
                    "conflict_reason": "; ".join(reasons),
                    "db_site": db_site,
                    "db_vlan_id": db_vlan_id,
                    "db_subnet": db_subnet,
                    "db_description": str(db_rec.get("description") or ""),
                }
                conflicts.append(conflict_rec)
            else:
                matched.append({
                    **rec,
                    "db_id": db_rec["id"],
                })

    return matched, pending_add, conflicts


# ── Output Writers ──────────────────────────────────────────────────────────

def write_csv(path: str, records: List[Dict[str, str]], fieldnames: List[str]) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)


MATCHED_FIELDS = [
    "ip_address", "subnet", "vlan_id", "vlan_name", "site",
    "description", "status", "db_id",
]
PENDING_FIELDS = [
    "ip_address", "subnet", "vlan_id", "vlan_name", "site",
    "description", "status",
]
CONFLICT_FIELDS = [
    "ip_address", "subnet", "vlan_id", "vlan_name", "site",
    "description", "status", "conflict_reason",
    "db_site", "db_vlan_id", "db_subnet", "db_description",
]


# ── Terminal Summary ───────────────────────────────────────────────────────

def print_summary(
    excel_count: int,
    matched: int,
    pending: int,
    conflicts: int,
    out_dir: str,
) -> None:
    w = 62
    sep = "─" * w
    print()
    print(sep)
    print(f"  NetBox Hub v4.3 — Offline IPAM Diff Report")
    print(sep)
    print(f"  Excel records parsed : {excel_count:>5}")
    print(f"  Matched (DB + sheet) : {matched:>5}")
    print(f"  Pending Add          : {pending:>5}")
    print(f"  Conflicts            : {conflicts:>5}")
    print(sep)
    print(f"  Output directory     : {out_dir}")
    print(f"    Matched  → {os.path.join(out_dir, 'ipam_matched.csv')}")
    print(f"    Pending  → {os.path.join(out_dir, 'ipam_pending_add.csv')}")
    print(f"    Conflicts→ {os.path.join(out_dir, 'ipam_conflicts.csv')}")
    print(sep)
    print()


# ── Main ────────────────────────────────────────────────────────────────────

def main() -> int:
    parser = argparse.ArgumentParser(
        description="NetBox Hub v4.3 — Offline IPAM Diff Engine",
    )
    parser.add_argument(
        "--excel",
        default="AW IPs & VLANs_Final (2).xlsx",
        help="Path to Excel inventory file",
    )
    parser.add_argument(
        "--db",
        default="data/sanitizer_vault.db",
        help="Path to local SQLite vault database",
    )
    parser.add_argument(
        "--out-dir",
        default="data/ipam_diff_output",
        help="Directory for output CSV reports",
    )
    args = parser.parse_args()

    excel_path = os.path.abspath(args.excel)
    db_path = os.path.abspath(args.db)
    out_dir = os.path.abspath(args.out_dir)

    if not os.path.isfile(excel_path):
        print(f"Error: Excel file not found: {excel_path}", file=sys.stderr)
        return 1

    if not os.path.isfile(db_path):
        print(f"Error: SQLite database not found: {db_path}", file=sys.stderr)
        return 1

    print(f"  Reading Excel: {excel_path}")
    parser_obj = ExcelIPAMParser(excel_path)
    excel_records = parser_obj.parse()
    print(f"  Parsed {len(excel_records)} valid IPAM records from Excel")

    print(f"  Connecting to DB: {db_path}")
    db = LocalSQLiteManager(db_path)
    db_records = db.get_all_records()
    print(f"  Found {len(db_records)} records in ipam_inventory")

    matched, pending, conflicts = classify_records(excel_records, db)

    write_csv(os.path.join(out_dir, "ipam_matched.csv"), matched, MATCHED_FIELDS)
    write_csv(os.path.join(out_dir, "ipam_pending_add.csv"), pending, PENDING_FIELDS)
    write_csv(os.path.join(out_dir, "ipam_conflicts.csv"), conflicts, CONFLICT_FIELDS)

    print_summary(len(excel_records), len(matched), len(pending), len(conflicts), out_dir)
    db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
